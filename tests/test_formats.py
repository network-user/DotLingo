from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest

from dotlingo.formats import (
    DocumentError,
    ScannedPdfError,
    export_document,
    import_document,
)


def test_txt_crlf_import_translation_export_round_trip(tmp_path: Path) -> None:
    source = tmp_path / "book.txt"
    original = "  First paragraph.\r\n\r\nSecond paragraph.  \n"
    source.write_bytes(original.encode("utf-8"))
    parsed = import_document(source)
    assert [block.text for block in parsed.blocks if block.translatable] == ["First paragraph.", "Second paragraph."]
    output = tmp_path / "translated.txt"
    export_document(source, output, parsed, {0: "Первый абзац.", 2: "Второй абзац."})
    assert output.read_bytes() == "  Первый абзац.\r\n\r\nВторой абзац.  \n".encode("utf-8")
    assert hashlib.sha256(source.read_bytes()).hexdigest() == hashlib.sha256(original.encode()).hexdigest()


def test_cp1251_import_is_supported(tmp_path: Path) -> None:
    source = tmp_path / "legacy.txt"
    source.write_bytes("Привет мир".encode("cp1251"))
    assert import_document(source).blocks[0].text == "Привет мир"


def test_markdown_export_keeps_heading_level_and_code_fence(tmp_path: Path) -> None:
    source = tmp_path / "notes.md"
    source.write_text("## Глава\n\nПервый абзац.\n\n```py\nx = 1\n```\n", encoding="utf-8")
    parsed = import_document(source)
    translations = {block.order: f"Перевод {block.order}" for block in parsed.blocks if block.translatable}
    output = tmp_path / "translated.md"
    export_document(source, output, parsed, translations)
    content = output.read_text(encoding="utf-8")
    assert content.startswith("## Перевод ")
    assert "```py\nx = 1\n```" in content


def test_docx_rejects_a_path_that_leaves_the_archive(tmp_path: Path) -> None:
    source = tmp_path / "escape.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("../evil.txt", "nope")
    with pytest.raises(DocumentError, match="небезопасный путь"):
        import_document(source)


def test_docx_import_and_export_preserves_table_and_does_not_mutate_source(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    source = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_heading("Chapter", level=1)
    document.add_paragraph("A paragraph.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left cell"
    table.cell(0, 1).text = "Right cell"
    document.save(source)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    parsed = import_document(source)
    assert any("смешанных фрагментов" in warning.lower() for warning in parsed.warnings)
    table_blocks = [block for block in parsed.blocks if "cell" in block.text.lower()]
    assert len(table_blocks) == 2
    translations = {block.order: f"T:{block.text}" for block in parsed.blocks if block.translatable}
    destination = tmp_path / "translated.docx"
    export_document(source, destination, parsed, translations)
    translated = docx.Document(destination)
    text = "\n".join(paragraph.text for paragraph in translated.paragraphs)
    table_text = " ".join(cell.text for row in translated.tables[0].rows for cell in row.cells)
    assert "T:Chapter" in text and "T:A paragraph." in text
    assert "T:Left cell" in table_text and "T:Right cell" in table_text
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest


def _make_epub(path: Path) -> None:
    container = '''<?xml version="1.0"?>
    <container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">
      <rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
    </container>'''
    package = '''<?xml version="1.0"?>
    <package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">
      <metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Fixture</dc:title></metadata>
      <manifest><item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/><item id="c2" href="c2.xhtml" media-type="application/xhtml+xml"/></manifest>
      <spine><itemref idref="c1"/><itemref idref="c2"/></spine>
    </package>'''
    first = '<html xmlns="http://www.w3.org/1999/xhtml"><body><h1>Chapter one</h1><p>Hello <em>world</em>.</p></body></html>'
    second = '<html xmlns="http://www.w3.org/1999/xhtml"><body><h1>Chapter two</h1><p>Second text.</p></body></html>'
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", package)
        archive.writestr("OEBPS/c1.xhtml", first)
        archive.writestr("OEBPS/c2.xhtml", second)
        archive.writestr("OEBPS/image.svg", "<svg />")


def test_epub_keeps_inline_code_breaks_and_parent_text(tmp_path: Path) -> None:
    source = tmp_path / "inline.epub"
    container = """<?xml version="1.0"?>
    <container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">
      <rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
    </container>"""
    package = """<?xml version="1.0"?>
    <package xmlns="http://www.idpf.org/2007/opf" version="3.0">
      <metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Inline</dc:title></metadata>
      <manifest><item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/></manifest>
      <spine><itemref idref="c1"/></spine>
    </package>"""
    chapter = (
        "<html xmlns=\"http://www.w3.org/1999/xhtml\"><body>"
        "<p>See <code>fig. 1</code> above.</p>"
        "<p>line1<br/>line2</p>"
        "<ul><li>Item <p>nested</p> tail</li></ul>"
        "</body></html>"
    )
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", package)
        archive.writestr("OEBPS/c1.xhtml", chapter)
    texts = [block.text for block in import_document(source).blocks]
    assert "See fig. 1 above." in texts
    assert any("line1" in item and "line2" in item and "\n" in item for item in texts)
    assert any(item.strip() == "nested" for item in texts)
    assert any("Item" in item and "tail" in item for item in texts)


def test_generic_exports_keep_the_text_and_the_source(tmp_path: Path) -> None:
    import hashlib
    import zipfile

    from docx import Document

    from dotlingo.formats import convert_document

    source = tmp_path / "note.txt"
    source.write_text("Глава\n\nПервый абзац про гавань.", encoding="utf-8")
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    expected = {
        ".html": ("<p>", "Первый абзац про гавань."),
        ".rtf": ("\\u", "rtf1"),
        ".fb2": ("<p>", "Первый абзац про гавань."),
        ".json": ('"text"', "Первый абзац про гавань."),
    }
    for suffix, needles in expected.items():
        destination = tmp_path / f"note{suffix}"
        convert_document(source, destination)
        text = destination.read_text(encoding="utf-8")
        for needle in needles:
            assert needle in text
    odt = tmp_path / "note.odt"
    convert_document(source, odt)
    with zipfile.ZipFile(odt) as archive:
        assert "Первый абзац про гавань." in archive.read("content.xml").decode("utf-8")
    docx_path = tmp_path / "note.docx"
    convert_document(source, docx_path)
    paragraphs = [item.text for item in Document(docx_path).paragraphs]
    assert any("Первый абзац про гавань." in item for item in paragraphs)
    epub_path = tmp_path / "note.epub"
    convert_document(source, epub_path)
    with zipfile.ZipFile(epub_path) as archive:
        assert archive.read("mimetype") == b"application/epub+zip"
        chapter = archive.read("OEBPS/chapter.xhtml").decode("utf-8")
    assert "Первый абзац про гавань." in chapter
    again = tmp_path / "note.txt"
    try:
        convert_document(source, again)
    except Exception as exc:
        assert "поверх" in str(exc)
    else:
        raise AssertionError("конвертер не должен затирать исходник")
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest


def test_plain_export_separates_blocks_without_their_own_breaks() -> None:
    from dotlingo.formats import Block, ParsedDocument, _text_export

    parsed = ParsedDocument(
        "Doc",
        "docx",
        (
            Block(0, "main", "Документ", "heading", "Chapter"),
            Block(1, "main", "Документ", "paragraph", "First paragraph."),
            Block(2, "main", "Документ", "paragraph", "Second paragraph."),
        ),
    )
    assert _text_export(parsed, {}, False) == "Chapter\n\nFirst paragraph.\n\nSecond paragraph."
    kept = ParsedDocument(
        "Doc",
        "txt",
        (
            Block(0, "main", "Документ", "paragraph", "One", suffix="\n\n"),
            Block(1, "main", "Документ", "paragraph", "Two"),
        ),
    )
    assert _text_export(kept, {}, False) == "One\n\nTwo"


def test_epub_round_trip_keeps_spine_order_and_non_text_assets(tmp_path: Path) -> None:
    source = tmp_path / "book.epub"
    _make_epub(source)
    parsed = import_document(source)
    headings = [block.text for block in parsed.blocks if block.kind == "heading"]
    assert headings == ["Chapter one", "Chapter two"]
    translations = {block.order: f"T:{block.text}" for block in parsed.blocks if block.translatable}
    output = tmp_path / "translated.epub"
    export_document(source, output, parsed, translations)
    with zipfile.ZipFile(output) as archive:
        assert archive.read("mimetype") == b"application/epub+zip"
        assert archive.read("OEBPS/image.svg") == b"<svg />"
        chapter1 = ET.fromstring(archive.read("OEBPS/c1.xhtml"))
        chapter2 = ET.fromstring(archive.read("OEBPS/c2.xhtml"))
        assert "T:Chapter one" in "".join(chapter1.itertext())
        assert "T:Hello world." in "".join(chapter1.itertext())
        assert "T:Chapter two" in "".join(chapter2.itertext())


def test_same_format_for_pdf_stays_pdf() -> None:
    from dotlingo.api import _resolve_output_suffix

    assert _resolve_output_suffix("pdf", False, "") == (".pdf", "")
    assert _resolve_output_suffix("pdf", True, "") == (".pdf", "")
    assert _resolve_output_suffix("pdf", False, "md") == (".md", "")


def test_pdf_text_layer_supported_and_scan_explicitly_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    pytest.importorskip("pypdf")
    fitz = pytest.importorskip("fitz")
    text_pdf = tmp_path / "text.pdf"
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "This is searchable text.")
        document.save(text_pdf)
    parsed = import_document(text_pdf)
    assert any("текстовый слой" in item.lower() for item in parsed.warnings)
    pytest.importorskip("reportlab")
    output = tmp_path / "translated.pdf"
    translations = {
        block.order: "Это искаемый текст."
        for block in parsed.blocks
        if block.translatable
    }
    export_document(text_pdf, output, parsed, translations)
    from pypdf import PdfReader

    text = "\n".join(page.extract_text() or "" for page in PdfReader(str(output)).pages)
    assert "Это искаемый текст." in text
    assert "Страница 1" not in text

    scan_pdf = tmp_path / "scan.pdf"
    with fitz.open() as document:
        document.new_page()
        document.save(scan_pdf)
    monkeypatch.setattr("dotlingo.image_layout.resolve_ocr_language", lambda requested=None: "en")
    monkeypatch.setattr("dotlingo.formats._recognize_page", lambda path, page, lang, scale: [])
    with pytest.raises(ScannedPdfError) as error:
        import_document(scan_pdf)
    assert error.value.pages == [1]
    assert "распознавание не нашло" in str(error.value)


def _dark_pixels(path: Path, box: tuple[int, int, int, int]) -> int:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(path))
    try:
        image = document[0].render(scale=2).to_pil().convert("RGB")
    finally:
        document.close()
    x0, y0, x1, y1 = box
    count = 0
    for y in range(y0, y1):
        for x in range(x0, x1):
            red, green, blue = image.getpixel((x, y))
            if red < 80 and green < 80 and blue < 80:
                count += 1
    return count


def test_scanned_pdf_paints_translation_on_the_page(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    pytest.importorskip("pypdfium2")
    pytest.importorskip("reportlab")
    from reportlab.pdfgen import canvas

    from dotlingo.image_layout import OcrLine

    source = tmp_path / "scan.pdf"
    writer = canvas.Canvas(str(source), pagesize=(200, 280))
    writer.showPage()
    writer.save()
    before = source.read_bytes()
    monkeypatch.setattr("dotlingo.image_layout.resolve_ocr_language", lambda requested=None: "en")
    monkeypatch.setattr(
        "dotlingo.formats._recognize_page",
        lambda path, page, lang, scale: [OcrLine(0, "Hello there", 20, 40, 180, 80)],
    )
    parsed = import_document(source, ocr_lang="en")
    assert parsed.metadata["pdfLayout"] == "scan"
    assert parsed.blocks[0].text == "Hello there"
    assert "вернётся PDF" in parsed.warnings[0]

    output = tmp_path / "scan.ru.pdf"
    export_document(source, output, parsed, {0: "Привет"})
    text_output = tmp_path / "scan.ru.txt"
    export_document(source, text_output, parsed, {0: "Привет"})
    assert source.read_bytes() == before
    assert "Привет" in text_output.read_text(encoding="utf-8")
    assert _dark_pixels(output, (20, 40, 180, 80)) > 10


def test_corrupt_pdf_is_rejected(tmp_path: Path) -> None:
    pytest.importorskip("pypdf")
    source = tmp_path / "broken.pdf"
    source.write_bytes(b"not a PDF")
    with pytest.raises(DocumentError, match="Не удалось прочитать PDF"):
        import_document(source)


def test_corrupt_epub_is_rejected_without_copying(tmp_path: Path) -> None:
    source = tmp_path / "broken.epub"
    source.write_bytes(b"not a zip")
    with pytest.raises(DocumentError):
        import_document(source)


def _write_line_pdf(path: Path, text: str, font: str) -> None:
    from reportlab.pdfgen import canvas

    sheet = canvas.Canvas(str(path), pagesize=(400, 300))
    sheet.setFont(font, 12)
    sheet.drawString(40, 200, text)
    sheet.save()


def _encrypt_pdf(source: Path, destination: Path, user_password: str, owner_password: str) -> None:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.append(PdfReader(str(source)))
    writer.encrypt(user_password, owner_password)
    with destination.open("wb") as handle:
        writer.write(handle)


def test_empty_user_password_pdf_imports(tmp_path: Path) -> None:
    pytest.importorskip("pypdf")
    pytest.importorskip("reportlab")
    plain = tmp_path / "plain.pdf"
    _write_line_pdf(plain, "This is searchable text for the lock.", "Helvetica")
    locked = tmp_path / "locked.pdf"
    _encrypt_pdf(plain, locked, "", "owner-secret")
    from pypdf import PdfReader

    assert PdfReader(str(locked)).is_encrypted
    parsed = import_document(locked)
    assert any("searchable text" in block.text for block in parsed.blocks)
    assert parsed.metadata.get("pdfLayout") != "book"


def test_real_pdf_password_is_rejected(tmp_path: Path) -> None:
    pytest.importorskip("pypdf")
    pytest.importorskip("reportlab")
    plain = tmp_path / "plain.pdf"
    _write_line_pdf(plain, "Closed page.", "Helvetica")
    locked = tmp_path / "locked.pdf"
    _encrypt_pdf(plain, locked, "secret", "owner-secret")
    with pytest.raises(DocumentError, match="паролем"):
        import_document(locked)


def test_book_pdf_with_empty_password_exports_pdf(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    pytest.importorskip("pypdf")
    pytest.importorskip("reportlab")
    pytest.importorskip("pdfminer")
    pytest.importorskip("pypdfium2")
    import dotlingo.pdf_layout as pdf_layout
    from dotlingo.pdf_layout import FONT_NAME, _register_font

    # В тестовом шрифте внутреннее имя OldStandard, не ModernMT. Правило полосы то же.
    monkeypatch.setattr(pdf_layout, "_BOOK_FONT", "OldStandard")
    _register_font()
    plain = tmp_path / "alice.pdf"
    _write_line_pdf(plain, "Alice was not hurt, and she jumped up.", FONT_NAME)
    locked = tmp_path / "alice-locked.pdf"
    _encrypt_pdf(plain, locked, "", "owner-secret")
    parsed = import_document(locked)
    assert parsed.metadata.get("pdfLayout") == "book"
    assert parsed.metadata.get("pageCount") == 1
    translations = {
        block.order: "Алиса совсем не ушиблась."
        for block in parsed.blocks
        if block.translatable
    }
    output = tmp_path / "alice-ru.pdf"
    export_document(locked, output, parsed, translations)
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(output))
    text = document[0].get_textpage().get_text_bounded()
    count = len(document)
    document.close()
    assert count == 1
    assert "Алиса совсем не ушиблась." in text


def test_book_pdf_that_does_not_fit_stays_pdf(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    pytest.importorskip("pypdf")
    pytest.importorskip("reportlab")
    pytest.importorskip("pdfminer")
    pytest.importorskip("pypdfium2")
    import dotlingo.pdf_layout as pdf_layout
    from dotlingo.pdf_layout import FONT_NAME, _register_font

    monkeypatch.setattr(pdf_layout, "_BOOK_FONT", "OldStandard")
    _register_font()
    plain = tmp_path / "alice.pdf"
    _write_line_pdf(plain, "Alice was not hurt, and she jumped up.", FONT_NAME)
    parsed = import_document(plain)
    assert parsed.metadata.get("pdfLayout") == "book"
    long = "Алиса " * 80
    translations = {block.order: long for block in parsed.blocks if block.translatable}
    output = tmp_path / "alice-long.pdf"
    export_document(plain, output, parsed, translations)
    assert output.suffix == ".pdf"
    assert output.is_file()
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(output))
    pages = [
        document[index].get_textpage().get_text_bounded()
        for index in range(len(document))
    ]
    count = len(document)
    document.close()
    assert count >= 2
    assert "Алиса" in "\n".join(pages)
    assert output.with_suffix(".md").exists() is False
