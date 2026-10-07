from __future__ import annotations

import json
import os
import posixpath
import re
import uuid
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree as ET

SUPPORTED_IMPORT = {".txt", ".md", ".markdown", ".docx", ".epub", ".pdf"}
SUPPORTED_EXPORT = {
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".rtf",
    ".fb2",
    ".odt",
    ".json",
    ".docx",
    ".epub",
    ".pdf",
}
MAX_DOCUMENT_BYTES = 512 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 20_000
MAX_ARCHIVE_UNPACKED = 512 * 1024 * 1024


class DocumentError(ValueError):
    """The selected file is damaged, unsupported, or cannot be converted safely."""


SCAN_SCALE = 2.0


class ScannedPdfError(DocumentError):
    def __init__(self, pages: list[int]) -> None:
        self.pages = pages
        page_list = ", ".join(map(str, pages[:12]))
        suffix = "…" if len(pages) > 12 else ""
        super().__init__(
            "На страницах "
            f"{page_list}{suffix} распознавание не нашло текста."
        )


@dataclass(frozen=True)
class Block:
    order: int
    section: str
    section_title: str
    kind: str
    text: str
    locator: dict[str, Any] = field(default_factory=dict)
    prefix: str = ""
    suffix: str = ""
    translatable: bool = True


@dataclass(frozen=True)
class ParsedDocument:
    title: str
    format: str
    blocks: tuple[Block, ...]
    warnings: tuple[str, ...] = ()
    metadata: dict[str, Any] = field(default_factory=dict)


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            return data.decode("cp1251")
        except UnicodeDecodeError as exc:
            raise DocumentError("Не удалось определить кодировку текста (поддерживаются UTF-8 и CP1251).") from exc


def _plain_blocks(text: str) -> tuple[Block, ...]:
    pieces = re.split(r"(\r?\n[\t ]*\r?\n+)", text)
    blocks: list[Block] = []
    for item in pieces:
        if not item:
            continue
        if item.isspace():
            blocks.append(Block(len(blocks), "main", "Документ", "gap", item, translatable=False))
            continue
        leading = item[: len(item) - len(item.lstrip())]
        trailing = item[len(item.rstrip()) :]
        body_end = len(item) - len(trailing) if trailing else len(item)
        body = item[len(leading) : body_end]
        if not body:
            blocks.append(Block(len(blocks), "main", "Документ", "gap", item, translatable=False))
        else:
            blocks.append(
                Block(len(blocks), "main", "Документ", "paragraph", body, prefix=leading, suffix=trailing)
            )
    return tuple(blocks)


def _markdown_blocks(text: str) -> tuple[Block, ...]:
    lines = text.splitlines(keepends=True)
    if text and not lines:
        lines = [text]
    blocks: list[Block] = []
    section_id, section_title, in_fence = "main", "Документ", False
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.lstrip()
        indent = line[: len(line) - len(stripped)]
        fence = re.match(r"^(```+|~~~+)", stripped)
        if fence:
            collected = [line]
            in_fence = not in_fence
            i += 1
            if in_fence:
                while i < len(lines):
                    collected.append(lines[i])
                    if re.match(r"^\s*(```+|~~~+)\s*$", lines[i]):
                        i += 1
                        in_fence = False
                        break
                    i += 1
            blocks.append(
                Block(len(blocks), section_id, section_title, "code", "".join(collected), translatable=False)
            )
            continue
        if in_fence:
            blocks.append(Block(len(blocks), section_id, section_title, "code", line, translatable=False))
            i += 1
            continue
        if not stripped.strip():
            blocks.append(Block(len(blocks), section_id, section_title, "gap", line, translatable=False))
            i += 1
            continue

        marker = re.match(r"^(#{1,6}\s+|>\s+|(?:[-+*]|\d+[.)])\s+)", stripped)
        prefix = indent + (marker.group(0) if marker else "")
        body_line = line[len(prefix) :]
        kind = "heading" if marker and marker.group(0).lstrip().startswith("#") else "paragraph"
        if kind == "heading":
            title = body_line.strip()
            section_id = f"section-{len([b for b in blocks if b.kind == 'heading']) + 1}"
            section_title = title or "Раздел"
        # Join wrapped Markdown lines into one prose block. Lists, quotes and headings
        # remain separate so their syntax can be restored exactly.
        if marker:
            content = body_line
            end = i + 1
        else:
            collected = [body_line]
            end = i + 1
            while end < len(lines) and lines[end].strip():
                next_stripped = lines[end].lstrip()
                if re.match(r"^(#{1,6}\s+|>\s+|(?:[-+*]|\d+[.)])\s+|```+|~~~+)", next_stripped):
                    break
                collected.append(lines[end])
                end += 1
            content = "".join(collected)
        trailing = content[len(content.rstrip("\r\n")) :]
        if trailing:
            content = content[: -len(trailing)]
        blocks.append(
            Block(
                len(blocks),
                section_id,
                section_title,
                kind,
                content,
                prefix=prefix,
                suffix=trailing,
            )
        )
        i = end
    return tuple(blocks)


def _safe_archive(zf: zipfile.ZipFile) -> None:
    infos = zf.infolist()
    if len(infos) > MAX_ARCHIVE_MEMBERS:
        raise DocumentError("В архиве слишком много элементов.")
    if sum(item.file_size for item in infos) > MAX_ARCHIVE_UNPACKED:
        raise DocumentError("Распакованный архив слишком велик для безопасного импорта.")
    names: set[str] = set()
    for info in infos:
        normalized = posixpath.normpath(info.filename.replace("\\", "/"))
        if normalized.startswith("../") or normalized.startswith("/") or "\x00" in info.filename:
            raise DocumentError("В архиве найден небезопасный путь.")
        if normalized in names:
            raise DocumentError("В архиве есть повторяющиеся имена файлов.")
        names.add(normalized)


def _docx_paragraphs(container: Any) -> list[Any]:
    from docx.document import Document as DocumentObject
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = container if isinstance(container, DocumentObject) else None
    parent = document.element.body if document else container._tc
    part = document.part if document else container.part
    paragraphs: list[Any] = []
    seen_cells: set[int] = set()
    for child in parent.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraphs.append(Paragraph(child, part))
        elif tag == "tbl":
            table = Table(child, part)
            for row in table.rows:
                for cell in row.cells:
                    cell_identity = id(cell._tc)
                    if cell_identity in seen_cells:
                        continue
                    seen_cells.add(cell_identity)
                    paragraphs.extend(_docx_paragraphs(cell))
    return paragraphs


def _import_docx(path: Path) -> ParsedDocument:
    from docx import Document

    doc = Document(path)
    blocks: list[Block] = []
    section_id, section_title = "main", "Документ"
    for index, paragraph in enumerate(_docx_paragraphs(doc)):
        text = paragraph.text
        style = paragraph.style.name if paragraph.style else ""
        kind = "heading" if style.lower().startswith("heading") or style.lower() == "title" else "paragraph"
        if kind == "heading" and text.strip():
            section_id, section_title = f"section-{index + 1}", text.strip()
        blocks.append(
            Block(
                len(blocks),
                section_id,
                section_title,
                kind,
                text,
                locator={"paragraph": index},
                translatable=bool(text.strip()),
            )
        )
    title = path.stem
    try:
        if doc.core_properties.title:
            title = doc.core_properties.title
    except Exception:
        pass
    return ParsedDocument(title, "docx", tuple(blocks), ("Сохраняются абзацы и таблицы; форматирование внутри смешанных фрагментов может упроститься.",))


_EPUB_TAGS = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "blockquote", "dt", "dd", "td", "th", "figcaption", "title", "div"}


def _tag(element: ET.Element) -> str:
    return element.tag.rsplit("}", 1)[-1].lower() if isinstance(element.tag, str) else ""


def _inline_text(element: ET.Element) -> str:
    """Текст абзаца: inline и <br> остаются, вложенный абзац идёт отдельным блоком."""
    tag = _tag(element)
    if tag in {"script", "style", "svg", "math"}:
        return ""
    if tag == "br":
        return "\n"
    parts: list[str] = [element.text or ""]
    for child in list(element):
        child_tag = _tag(child)
        if child_tag in _EPUB_TAGS:
            pass
        elif child_tag == "br":
            parts.append("\n")
        elif child_tag not in {"script", "style", "svg", "math"}:
            parts.append(_inline_text(child))
        parts.append(child.tail or "")
    return "".join(parts)


def _epub_text_elements(root: ET.Element) -> list[ET.Element]:
    return [
        element
        for element in root.iter()
        if _tag(element) in _EPUB_TAGS and _inline_text(element).strip()
    ]


def _epub_package(zf: zipfile.ZipFile) -> tuple[str, str, list[tuple[str, str]]]:
    try:
        container = ET.fromstring(zf.read("META-INF/container.xml"))
        rootfile = next(item for item in container.iter() if _tag(item) == "rootfile")
        package_path = rootfile.attrib["full-path"]
        package = ET.fromstring(zf.read(package_path))
    except (KeyError, ET.ParseError, StopIteration, AttributeError) as exc:
        raise DocumentError("Повреждённый EPUB: отсутствует container.xml или package document.") from exc
    manifest = {}
    for item in package.iter():
        if _tag(item) == "item" and item.get("id") and item.get("href"):
            href = unquote(urlsplit(item.get("href", "")).path)
            resolved = posixpath.normpath(posixpath.join(posixpath.dirname(package_path), href))
            if resolved.startswith("../") or resolved.startswith("/"):
                raise DocumentError("В EPUB найден путь за пределами архива.")
            manifest[item.get("id", "")] = (resolved, item.get("media-type", ""))
    spine: list[tuple[str, str]] = []
    for itemref in package.iter():
        if _tag(itemref) == "itemref":
            entry = manifest.get(itemref.get("idref", ""))
            if entry and entry[1] in {"application/xhtml+xml", "text/html"}:
                spine.append(entry)
    if not spine:
        raise DocumentError("В EPUB не найдены XHTML-разделы в порядке книги.")
    title = next((item.text for item in package.iter() if _tag(item) == "title" and item.text), "")
    return package_path, title, spine


def _import_epub(path: Path) -> ParsedDocument:
    blocks: list[Block] = []
    warnings: list[str] = []
    try:
        with zipfile.ZipFile(path) as zf:
            _safe_archive(zf)
            _, title, spine = _epub_package(zf)
            for chapter_no, (href, _) in enumerate(spine, 1):
                try:
                    root = ET.fromstring(zf.read(href))
                except (KeyError, ET.ParseError) as exc:
                    raise DocumentError(f"Повреждён раздел EPUB: {href}") from exc
                current_section, current_title = f"chapter-{chapter_no}", f"Глава {chapter_no}"
                for item_index, element in enumerate(_epub_text_elements(root)):
                    text = _inline_text(element)
                    tag = _tag(element)
                    kind = "heading" if tag.startswith("h") and len(tag) == 2 else "paragraph"
                    if kind == "heading" and text.strip():
                        current_section, current_title = f"chapter-{chapter_no}-section-{item_index}", text.strip()
                    blocks.append(
                        Block(
                            len(blocks),
                            current_section,
                            current_title,
                            kind,
                            text,
                            locator={"href": href, "element": item_index},
                            translatable=bool(text.strip()),
                        )
                    )
    except (OSError, zipfile.BadZipFile) as exc:
        raise DocumentError("Файл EPUB повреждён или не является ZIP-архивом.") from exc
    warnings.append("Порядок глав, XHTML-структура, ссылки и изображения сохраняются; форматирование внутри абзаца может упроститься.")
    return ParsedDocument(title or path.stem, "epub", tuple(blocks), tuple(warnings))


def _open_pdf_reader(path: Path):
    """Открыть PDF. Пустой пароль пользователя снимается: так закрыт только доступ владельца."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentError("Импорт PDF требует дополнительной зависимости pypdf. Установите пакет pdf.") from exc
    try:
        doc = PdfReader(str(path), strict=True)
    except Exception as exc:
        raise DocumentError("Не удалось прочитать PDF. Проверьте, что файл не повреждён или защищён паролем.") from exc
    if doc.is_encrypted:
        try:
            opened = doc.decrypt("")
        except Exception as exc:
            raise DocumentError("PDF защищён паролем. Снимите защиту перед импортом.") from exc
        if not opened:
            raise DocumentError("PDF защищён паролем. Снимите защиту перед импортом.")
    return doc


def _book_layout_document(path: Path, page_count: int) -> ParsedDocument | None:
    """Абзацы книжной полосы. None, если это не такая книга или укладка недоступна."""
    try:
        from dotlingo.pdf_layout import extract_paragraphs
    except ImportError:
        return None
    try:
        paragraphs = extract_paragraphs(path, list(range(1, page_count + 1)), "book")
    except Exception:
        return None
    if not paragraphs:
        return None
    blocks = [
        Block(
            order,
            f"page-{item.page}",
            f"Страница {item.page}",
            "heading" if item.kind == "header" else "paragraph",
            item.source,
            locator={"page": item.page, "index": item.index, "layout": "book"},
        )
        for order, item in enumerate(paragraphs)
    ]
    return ParsedDocument(
        path.stem,
        "pdf",
        tuple(blocks),
        ("Книга вернётся PDF: текст заменяется, картинки и номера страниц остаются.",),
        {"pdfLayout": "book", "pageCount": page_count},
    )


def _recognize_page(path: Path, page: int, lang: str, scale: float) -> list[Any]:
    """Строки скана. Вынесено отдельно, чтобы тест подменял распознавание."""
    from dotlingo.image_layout import OcrUnavailable, ocr_image, render_pdf_page

    try:
        image, _width, _height = render_pdf_page(path, page, scale)
        return ocr_image(image, lang)
    except OcrUnavailable:
        raise
    except ImportError as exc:
        raise OcrUnavailable("Для сканов нужны пакеты pdf и winocr.") from exc


def _scan_locator(page_no: int, lines: list[Any]) -> dict[str, Any]:
    return {
        "page": page_no,
        "layout": "scan",
        "scale": SCAN_SCALE,
        "lines": [
            {"x0": line.x0, "y0": line.y0, "x1": line.x1, "y1": line.y1}
            for line in lines
        ],
    }


def _lines_from_locator(locator: dict[str, Any]) -> list[Any]:
    from dotlingo.image_layout import OcrLine

    lines: list[OcrLine] = []
    for item in locator.get("lines") or []:
        lines.append(
            OcrLine(
                len(lines),
                "",
                int(item["x0"]),
                int(item["y0"]),
                int(item["x1"]),
                int(item["y1"]),
            )
        )
    return lines


def _import_scan_pdf(
    path: Path,
    page_texts: list[str],
    image_only: list[int],
    ocr_lang: str | None,
) -> ParsedDocument:
    from dotlingo.image_layout import (
        OcrUnavailable,
        group_ocr_lines,
        page_sizes,
        paragraph_text,
        resolve_ocr_language,
        text_layer_lines,
    )

    try:
        lang = resolve_ocr_language(ocr_lang)
        sizes = page_sizes(path)
    except OcrUnavailable as exc:
        raise DocumentError(str(exc)) from exc
    except ImportError as exc:
        raise DocumentError("Для сканов нужны пакеты pdf и winocr.") from exc
    except Exception as exc:
        raise DocumentError("Не удалось открыть страницы PDF для распознавания.") from exc
    if len(sizes) != len(page_texts):
        raise DocumentError("Не удалось открыть страницы PDF для распознавания.")

    scan_pages = set(image_only)
    blocks: list[Block] = []
    recognized: list[int] = []
    failed: list[int] = []
    for page_no, _content in enumerate(page_texts, 1):
        try:
            if page_no in scan_pages:
                lines = _recognize_page(path, page_no, lang, SCAN_SCALE)
            else:
                _width, height = sizes[page_no - 1]
                lines = text_layer_lines(path, page_no, SCAN_SCALE, height)
        except OcrUnavailable as exc:
            raise DocumentError(str(exc)) from exc
        except Exception as exc:
            raise DocumentError(f"Не удалось прочитать страницу {page_no}.") from exc
        page_blocks: list[Block] = []
        for group in group_ocr_lines(lines):
            text = paragraph_text(group)
            if not text:
                continue
            digits = "".join(text.split()).isdigit()
            page_blocks.append(
                Block(
                    0,
                    f"page-{page_no}",
                    f"Страница {page_no}",
                    "paragraph",
                    text,
                    locator=_scan_locator(page_no, group),
                    translatable=not digits,
                )
            )
        if page_no in scan_pages:
            if page_blocks:
                recognized.append(page_no)
            else:
                failed.append(page_no)
        blocks.extend(page_blocks)
    if not recognized:
        raise ScannedPdfError(failed or image_only)
    ordered = [
        Block(
            index,
            block.section,
            block.section_title,
            block.kind,
            block.text,
            block.locator,
            block.prefix,
            block.suffix,
            block.translatable,
        )
        for index, block in enumerate(blocks)
    ]
    warning = (
        "Скан вернётся PDF: перевод пишется поверх страницы, картинка остаётся. "
        "Проверьте распознанный текст перед переводом."
    )
    if failed:
        missed = ", ".join(map(str, failed[:12]))
        tail = "…" if len(failed) > 12 else ""
        warning = f"{warning} Пустые страницы: {missed}{tail}."
    return ParsedDocument(
        path.stem,
        "pdf",
        tuple(ordered),
        (warning,),
        {
            "pdfLayout": "scan",
            "pageCount": len(page_texts),
            "ocrScale": SCAN_SCALE,
            "ocrLang": lang,
        },
    )


def _import_pdf(path: Path, ocr_lang: str | None = None) -> ParsedDocument:
    doc = _open_pdf_reader(path)
    try:
        page_count = len(doc.pages)
    except Exception as exc:
        raise DocumentError("Не удалось прочитать PDF. Проверьте, что файл не повреждён или защищён паролем.") from exc
    if page_count < 1:
        raise DocumentError("PDF не содержит страниц.")
    laid_out = _book_layout_document(path, page_count)
    if laid_out is not None:
        return laid_out
    try:
        page_texts = [page.extract_text() or "" for page in doc.pages]
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("Не удалось прочитать PDF. Проверьте, что файл не повреждён или защищён паролем.") from exc
    if not page_texts:
        raise DocumentError("PDF не содержит страниц.")
    image_only = [i + 1 for i, content in enumerate(page_texts) if len(content.strip()) < 12]
    if image_only:
        return _import_scan_pdf(path, page_texts, image_only, ocr_lang)
    blocks: list[Block] = []
    for page_no, content in enumerate(page_texts, 1):
        blocks.append(
            Block(
                len(blocks),
                f"page-{page_no}",
                f"Страница {page_no}",
                "heading",
                f"Страница {page_no}",
                translatable=False,
            )
        )
        for block in _plain_blocks(content):
            blocks.append(
                Block(
                    len(blocks),
                    f"page-{page_no}",
                    f"Страница {page_no}",
                    block.kind,
                    block.text,
                    prefix=block.prefix,
                    suffix=block.suffix,
                    translatable=block.translatable,
                )
            )
    warning = (
        "Из PDF извлекается только текстовый слой. "
        "Вёрстка, таблицы и изображения не переносятся. Перевод собирается новым PDF."
    )
    return ParsedDocument(path.stem, "pdf", tuple(blocks), (warning,))


def import_document(path: Path, ocr_lang: str | None = None) -> ParsedDocument:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_IMPORT:
        raise DocumentError(f"Формат {suffix or '(без расширения)'} не поддерживается для импорта.")
    try:
        if path.stat().st_size > MAX_DOCUMENT_BYTES:
            raise DocumentError("Размер документа больше 512 МБ.")
        if suffix in {".txt", ".md", ".markdown"}:
            text = _decode_text(path.read_bytes())
            blocks = _markdown_blocks(text) if suffix != ".txt" else _plain_blocks(text)
            return ParsedDocument(path.stem, "markdown" if suffix != ".txt" else "txt", blocks)
        if suffix == ".docx":
            return _import_docx(path)
        if suffix == ".epub":
            return _import_epub(path)
        if suffix == ".pdf":
            return _import_pdf(path, ocr_lang)
    except OSError as exc:
        raise DocumentError("Не удалось прочитать выбранный файл.") from exc
    raise DocumentError("Неподдерживаемый формат.")


def block_to_dict(block: Block) -> dict[str, Any]:
    return asdict(block)


def block_from_dict(data: dict[str, Any]) -> Block:
    return Block(**data)


def _translation_for(block: Block, translations: dict[int, str]) -> str:
    if not block.translatable:
        return block.text
    return translations.get(block.order, block.text)


def _text_export(parsed: ParsedDocument, translations: dict[int, str], markdown: bool) -> str:
    chunks: list[str] = []
    previous_section = None
    for block in parsed.blocks:
        if _is_page_marker(block):
            continue
        if markdown and block.section != previous_section and block.section_title and block.kind != "heading":
            if chunks and chunks[-1] and not chunks[-1].endswith("\n\n"):
                chunks.append("\n\n")
            chunks.append(f"## {block.section_title}\n\n")
        previous_section = block.section
        if block.kind == "heading" and markdown:
            body = _translation_for(block, translations).strip()
            heading_prefix = block.prefix if block.prefix.lstrip().startswith("#") else "# "
            chunks.append(f"{heading_prefix}{body}{block.suffix or chr(10)}")
        else:
            piece = block.prefix + _translation_for(block, translations) + block.suffix
            if chunks and piece and not chunks[-1][-1].isspace() and not piece[0].isspace():
                chunks.append("\n\n")
            chunks.append(piece)
    return "".join(chunks)


def _set_xml_element_text(element: ET.Element, value: str) -> None:
    element.text = value
    for child in element.iter():
        if child is element:
            continue
        child.text = None
        child.tail = None


def _export_epub(source: Path, destination: Path, parsed: ParsedDocument, translations: dict[int, str]) -> None:
    grouped: dict[str, dict[int, str]] = {}
    for block in parsed.blocks:
        if block.translatable and "href" in block.locator:
            grouped.setdefault(str(block.locator["href"]), {})[int(block.locator["element"])] = _translation_for(block, translations)
    try:
        with zipfile.ZipFile(source) as src, zipfile.ZipFile(destination, "w") as dst:
            _safe_archive(src)
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename in grouped:
                    root = ET.fromstring(data)
                    elements = _epub_text_elements(root)
                    for index, translated in grouped[item.filename].items():
                        if index < len(elements):
                            _set_xml_element_text(elements[index], translated)
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                clone = zipfile.ZipInfo(item.filename, item.date_time)
                clone.compress_type = zipfile.ZIP_STORED if item.filename == "mimetype" else item.compress_type
                clone.external_attr = item.external_attr
                clone.comment = item.comment
                dst.writestr(clone, data)
    except (OSError, zipfile.BadZipFile, ET.ParseError) as exc:
        raise DocumentError("Не удалось записать EPUB. Исходник сохранён без изменений.") from exc


def _export_docx(source: Path, destination: Path, parsed: ParsedDocument, translations: dict[int, str]) -> None:
    from docx import Document

    doc = Document(source)
    paragraphs = _docx_paragraphs(doc)
    for block in parsed.blocks:
        if not block.translatable:
            continue
        index = int(block.locator.get("paragraph", -1))
        if index < 0 or index >= len(paragraphs):
            continue
        value = _translation_for(block, translations)
        paragraph = paragraphs[index]
        text_nodes = paragraph._p.xpath(".//w:t")
        if text_nodes:
            text_nodes[0].text = value
            for node in text_nodes[1:]:
                node.text = ""
        elif value:
            paragraph.add_run(value)
    doc.save(destination)


def _layout_translations(parsed: ParsedDocument, translations: dict[int, str]) -> dict[tuple[int, int], str]:
    placed: dict[tuple[int, int], str] = {}
    for block in parsed.blocks:
        locator = block.locator
        if locator.get("layout") != "book":
            continue
        placed[(int(locator["page"]), int(locator["index"]))] = _translation_for(block, translations)
    return placed


_PAGE_MARKER = re.compile(r"Страница \d+")


def _is_page_marker(block: Block) -> bool:
    """Служебный заголовок страницы, который импорт добавляет к обычному PDF."""
    return (
        block.kind == "heading"
        and not block.locator
        and _PAGE_MARKER.fullmatch(block.text.strip()) is not None
    )


def _pdf_markup(text: str) -> str:
    from xml.sax.saxutils import escape

    return escape(text).replace("\n", "<br/>")


def _export_pdf_text(
    destination: Path,
    parsed: ParsedDocument,
    translations: dict[int, str],
) -> None:
    """Новый PDF с текстом перевода. Полоса оригинала здесь не копируется."""
    try:
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise DocumentError("Для сборки PDF не хватает библиотеки reportlab.") from exc
    try:
        from dotlingo.pdf_layout import FONT_NAME, LayoutError, _register_font

        _register_font()
    except ImportError as exc:
        raise DocumentError("Для сборки PDF не хватает библиотек укладки.") from exc
    except LayoutError as exc:
        raise DocumentError(str(exc)) from exc

    body = ParagraphStyle(
        "dotlingo-body",
        fontName=FONT_NAME,
        fontSize=11,
        leading=15,
        alignment=TA_LEFT,
    )
    heading = ParagraphStyle(
        "dotlingo-heading",
        parent=body,
        fontSize=14,
        leading=18,
        spaceBefore=10,
        spaceAfter=4,
    )
    sections: list[list[Block]] = []
    current: str | None = None
    group: list[Block] = []
    for block in parsed.blocks:
        if _is_page_marker(block):
            continue
        if block.section != current:
            if group:
                sections.append(group)
            group = []
            current = block.section
        group.append(block)
    if group:
        sections.append(group)

    story: list[Any] = []
    for index, blocks in enumerate(sections):
        if index:
            story.append(PageBreak())
        for block in blocks:
            raw = f"{block.prefix}{_translation_for(block, translations)}{block.suffix}".strip()
            if not raw:
                continue
            story.append(Paragraph(_pdf_markup(raw), heading if block.kind == "heading" else body))
            story.append(Spacer(1, 8))
    if not story:
        story.append(Paragraph(" ", body))

    document = SimpleDocTemplate(
        str(destination),
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=56,
        bottomMargin=56,
        title=(parsed.title or "")[:120],
    )
    try:
        document.build(story)
    except Exception as exc:
        raise DocumentError("Не удалось собрать PDF.") from exc


def _draw_scan_overflow(writer: Any, text: str) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase.pdfmetrics import stringWidth

    from dotlingo.pdf_layout import FONT_NAME, LayoutError, _register_font

    try:
        _register_font()
    except LayoutError as exc:
        raise DocumentError(str(exc)) from exc
    width, height = A4
    font_size = 11
    margin = 54
    leading = 15
    max_width = width - (margin * 2)

    def new_page() -> float:
        writer.showPage()
        writer.setPageSize((width, height))
        writer.setFont(FONT_NAME, font_size)
        return height - 56

    y = new_page()
    for paragraph in text.split("\n\n"):
        line = ""
        for word in paragraph.split():
            trial = word if not line else f"{line} {word}"
            if stringWidth(trial, FONT_NAME, font_size) <= max_width:
                line = trial
                continue
            if line:
                if y < 56:
                    y = new_page()
                writer.drawString(margin, y, line)
                y -= leading
            line = word
        if line:
            if y < 56:
                y = new_page()
            writer.drawString(margin, y, line)
            y -= leading
        y -= leading


def _export_pdf_scan(
    source_path: Path,
    destination: Path,
    parsed: ParsedDocument,
    translations: dict[int, str],
) -> None:
    """Картинка каждой страницы. Перевод пишется в рамки распознанных строк."""
    if parsed.metadata.get("pdfLayout") != "scan":
        raise DocumentError("Для этого PDF нет скана.")
    try:
        page_count = int(parsed.metadata.get("pageCount") or 0)
        scale = float(parsed.metadata.get("ocrScale") or SCAN_SCALE)
    except (TypeError, ValueError) as exc:
        raise DocumentError("В PDF не записаны страницы скана.") from exc
    if page_count < 1:
        raise DocumentError("В PDF не записаны страницы скана.")
    try:
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise DocumentError("Для сборки PDF не хватает библиотеки reportlab.") from exc
    try:
        from dotlingo.image_layout import OcrUnavailable, place_scan_groups, render_pdf_page
    except ImportError as exc:
        raise DocumentError("Для сборки скана не хватает библиотек pdf.") from exc

    by_page: dict[int, list[tuple[list[Any], str]]] = {}
    for block in parsed.blocks:
        locator = block.locator
        if locator.get("layout") != "scan":
            continue
        try:
            page = int(locator.get("page") or 0)
        except (TypeError, ValueError):
            continue
        lines = _lines_from_locator(locator)
        if not lines:
            continue
        translated = _translation_for(block, translations).strip()
        if not block.translatable or not translated or translated == block.text.strip():
            continue
        by_page.setdefault(page, []).append((lines, translated))

    writer = None
    overflow: list[str] = []
    for page in range(1, page_count + 1):
        try:
            image, width, height = render_pdf_page(source_path, page, scale)
        except Exception as exc:
            raise DocumentError(f"Не удалось открыть страницу {page} для сборки PDF.") from exc
        groups = [lines for lines, _text in by_page.get(page, [])]
        placed = {index: text for index, (_lines, text) in enumerate(by_page.get(page, []))}
        try:
            if groups:
                image, leftover = place_scan_groups(image, groups, placed)
                if leftover:
                    overflow.append(leftover)
        except OcrUnavailable as exc:
            raise DocumentError(str(exc)) from exc
        if writer is None:
            writer = canvas.Canvas(str(destination), pagesize=(width, height))
        else:
            writer.showPage()
            writer.setPageSize((width, height))
        writer.drawImage(ImageReader(image), 0, 0, width=width, height=height, mask="auto")
    if writer is None:
        raise DocumentError("PDF не содержит страниц.")
    if overflow:
        _draw_scan_overflow(writer, "\n\n".join(overflow))
    writer.save()


def _export_pdf_layout(
    source_path: Path,
    destination: Path,
    parsed: ParsedDocument,
    translations: dict[int, str],
) -> None:
    if parsed.metadata.get("pdfLayout") != "book":
        raise DocumentError("Для этого PDF нет книжной полосы.")
    try:
        page_count = int(parsed.metadata.get("pageCount") or 0)
    except (TypeError, ValueError) as exc:
        raise DocumentError("В PDF не записано число страниц.") from exc
    if page_count < 1:
        raise DocumentError("В PDF не записано число страниц.")
    try:
        from dotlingo.pdf_layout import LayoutError, place_translations
    except ImportError as exc:
        raise DocumentError("Для сборки PDF не хватает библиотек укладки.") from exc
    try:
        place_translations(
            source_path,
            destination,
            _layout_translations(parsed, translations),
            pages=list(range(1, page_count + 1)),
            scope="book",
        )
    except LayoutError as exc:
        raise DocumentError(str(exc)) from exc


def _export_pieces(parsed: ParsedDocument, translations: dict[int, str]) -> list[tuple[str, str]]:
    """Абзацы для новой вёрстки. Служебные метки страниц и пустые куски не входят."""
    pieces: list[tuple[str, str]] = []
    for block in parsed.blocks:
        if _is_page_marker(block) or block.kind == "gap":
            continue
        text = _translation_for(block, translations).strip()
        if not text:
            continue
        pieces.append(("heading" if block.kind == "heading" else "paragraph", text))
    return pieces


def _xml_text(value: str) -> str:
    return (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _html_export(parsed: ParsedDocument, translations: dict[int, str]) -> str:
    title = _xml_text(parsed.title or "Документ")
    parts = [
        "<!DOCTYPE html>",
        '<html lang="und">',
        "<head>",
        '<meta charset="utf-8">',
        f"<title>{title}</title>",
        "</head>",
        "<body>",
        f"<article><h1>{title}</h1>",
    ]
    for kind, text in _export_pieces(parsed, translations):
        body = "<br>\n".join(_xml_text(line) for line in text.splitlines()) or "<br>"
        tag = "h2" if kind == "heading" else "p"
        parts.append(f"<{tag}>{body}</{tag}>")
    parts.extend(["</article>", "</body>", "</html>", ""])
    return "\n".join(parts)


def _rtf_escape(text: str) -> str:
    chunks: list[str] = []
    for char in text.replace("\r\n", "\n").replace("\r", "\n"):
        if char == "\n":
            chunks.append(r"\line ")
            continue
        if char in "\\{}":
            chunks.append("\\" + char)
            continue
        code = ord(char)
        if code < 128:
            chunks.append(char)
            continue
        if code > 32767:
            code -= 65536
        chunks.append(f"\\u{code}?")
    return "".join(chunks)


def _rtf_export(parsed: ParsedDocument, translations: dict[int, str]) -> str:
    body = [r"{\rtf1\ansi\ansicpg65001\uc1\deff0{\fonttbl{\f0\fnil Segoe UI;}}\f0\fs22"]
    title = (parsed.title or "").strip()
    if title:
        body.append(r"\fs32\b " + _rtf_escape(title) + r"\b0\fs22\par ")
    for kind, text in _export_pieces(parsed, translations):
        if kind == "heading":
            body.append(r"\fs28\b " + _rtf_escape(text) + r"\b0\fs22\par ")
        else:
            body.append(_rtf_escape(text) + r"\par ")
    body.append("}")
    return "".join(body)


def _fb2_export(parsed: ParsedDocument, translations: dict[int, str]) -> str:
    title = _xml_text(parsed.title or "Документ")
    sections: list[str] = []
    current: list[str] = []
    heading = ""
    for kind, text in _export_pieces(parsed, translations):
        if kind == "heading":
            if current or heading:
                sections.append(_fb2_section(heading or title, current))
            heading = text
            current = []
            continue
        current.append(text)
    sections.append(_fb2_section(heading or title, current))
    body = "".join(sections)
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0">\n'
        "<description><title-info>\n"
        f"<book-title>{title}</book-title>\n"
        "<lang>und</lang>\n"
        "</title-info></description>\n"
        f"<body>{body}</body>\n"
        "</FictionBook>\n"
    )


def _fb2_section(title: str, paragraphs: list[str]) -> str:
    lines = ["<section>", f"<title><p>{_xml_text(title)}</p></title>"]
    lines.extend(f"<p>{_xml_text(item)}</p>" for item in paragraphs)
    if not paragraphs:
        lines.append("<p></p>")
    lines.append("</section>")
    return "".join(lines)


def _odt_export(destination: Path, parsed: ParsedDocument, translations: dict[int, str]) -> None:
    title = _xml_text(parsed.title or "Документ")
    nodes = [f'<text:h text:outline-level="1">{title}</text:h>']
    for kind, text in _export_pieces(parsed, translations):
        tag = "text:h" if kind == "heading" else "text:p"
        level = ' text:outline-level="2"' if kind == "heading" else ""
        lines = text.splitlines() or [""]
        inner = "<text:line-break/>".join(_xml_text(line) for line in lines)
        nodes.append(f"<{tag}{level}>{inner}</{tag}>")
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
        f'<office:body><office:text>{"".join(nodes)}</office:text></office:body>'
        "</office:document-content>"
    )
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
        'manifest:version="1.2">'
        '<manifest:file-entry manifest:full-path="/" '
        'manifest:media-type="application/vnd.oasis.opendocument.text"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>"
    )
    with zipfile.ZipFile(destination, "w") as archive:
        info = zipfile.ZipInfo("mimetype")
        info.compress_type = zipfile.ZIP_STORED
        archive.writestr(info, "application/vnd.oasis.opendocument.text")
        archive.writestr("META-INF/manifest.xml", manifest)
        archive.writestr("content.xml", content)


def _json_export(parsed: ParsedDocument, translations: dict[int, str]) -> str:
    return json.dumps(
        {
            "title": parsed.title,
            "sourceFormat": parsed.format,
            "blocks": [
                {"kind": kind, "text": text} for kind, text in _export_pieces(parsed, translations)
            ],
        },
        ensure_ascii=False,
        indent=2,
    )


def _export_docx_plain(destination: Path, parsed: ParsedDocument, translations: dict[int, str]) -> None:
    from docx import Document

    document = Document()
    title = (parsed.title or "").strip()
    if title:
        document.core_properties.title = title[:200]
        document.add_heading(title[:400], level=0)
    for kind, text in _export_pieces(parsed, translations):
        lines = text.splitlines() or [""]
        if kind == "heading":
            document.add_heading(lines[0][:400], level=1)
            lines = lines[1:]
            if not lines:
                continue
        paragraph = document.add_paragraph()
        for index, line in enumerate(lines):
            if index:
                paragraph.add_run().add_break()
            paragraph.add_run(line)
    document.save(destination)


def _export_epub_plain(destination: Path, parsed: ParsedDocument, translations: dict[int, str]) -> None:
    title = _xml_text(parsed.title or "Документ")
    body = [f"<h1>{title}</h1>"]
    for kind, text in _export_pieces(parsed, translations):
        inner = "<br/>".join(_xml_text(line) for line in text.splitlines())
        tag = "h2" if kind == "heading" else "p"
        body.append(f"<{tag}>{inner}</{tag}>")
    chapter = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<html xmlns="http://www.w3.org/1999/xhtml"><head>'
        f"<title>{title}</title></head><body>{''.join(body)}</body></html>"
    )
    container = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" '
        'media-type="application/oebps-package+xml"/></rootfiles></container>'
    )
    package = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
        "<metadata xmlns:dc=\"http://purl.org/dc/elements/1.1/\">"
        f"<dc:title>{title}</dc:title><dc:language>und</dc:language>"
        '<dc:identifier id="id">dotlingo-converted</dc:identifier></metadata>'
        '<manifest><item id="c1" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest>'
        "<spine><itemref idref=\"c1\"/></spine></package>"
    )
    with zipfile.ZipFile(destination, "w") as archive:
        info = zipfile.ZipInfo("mimetype")
        info.compress_type = zipfile.ZIP_STORED
        archive.writestr(info, "application/epub+zip")
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", package)
        archive.writestr("OEBPS/chapter.xhtml", chapter)


def converted_destination(source: Path, suffix: str, folder: Path | None = None) -> Path:
    """Новый путь рядом с файлом или в выбранной папке. Исходное имя не занимается."""
    directory = folder if folder is not None else source.parent
    stem = source.stem or "document"
    candidate = directory / f"{stem}{suffix}"
    try:
        same = candidate.resolve() == source.resolve()
    except OSError:
        same = False
    if same or candidate.exists():
        candidate = directory / f"{stem}.converted{suffix}"
    index = 2
    while candidate.exists():
        candidate = directory / f"{stem}.converted-{index}{suffix}"
        index += 1
    return candidate


def convert_document(source: Path, destination: Path) -> None:
    """Собрать другой формат из текста файла. Исходный файл не открывается на запись."""
    source = Path(source)
    destination = Path(destination)
    try:
        if destination.resolve() == source.resolve():
            raise DocumentError("Нельзя записать результат поверх исходного файла.")
    except OSError as exc:
        raise DocumentError("Не удалось сравнить пути исходного файла и результата.") from exc
    before = source.read_bytes()
    parsed = import_document(source)
    export_document(source, destination, parsed, {})
    if source.read_bytes() != before:
        raise DocumentError("Исходный файл изменился во время конвертации.")


def export_document(
    source_path: Path,
    destination: Path,
    parsed: ParsedDocument,
    translations: dict[int, str],
) -> None:
    """Write a separate translated artifact; never modify the saved source copy."""
    destination = Path(destination)
    suffix = destination.suffix.lower()
    if suffix not in SUPPORTED_EXPORT:
        raise DocumentError("Для этого формата нет проверенного экспортёра.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(f".{destination.stem}.{uuid.uuid4().hex}.partial{suffix}")
    try:
        if suffix in {".txt", ".md", ".markdown"}:
            content = _text_export(parsed, translations, suffix != ".txt")
            temporary.write_text(content, encoding="utf-8", newline="")
        elif suffix == ".html":
            temporary.write_text(_html_export(parsed, translations), encoding="utf-8", newline="\n")
        elif suffix == ".rtf":
            temporary.write_text(_rtf_export(parsed, translations), encoding="utf-8", newline="\n")
        elif suffix == ".fb2":
            temporary.write_text(_fb2_export(parsed, translations), encoding="utf-8", newline="\n")
        elif suffix == ".json":
            temporary.write_text(_json_export(parsed, translations), encoding="utf-8", newline="\n")
        elif suffix == ".odt":
            _odt_export(temporary, parsed, translations)
        elif suffix == ".docx":
            if parsed.format == "docx":
                _export_docx(source_path, temporary, parsed, translations)
            else:
                _export_docx_plain(temporary, parsed, translations)
        elif suffix == ".epub":
            if parsed.format == "epub":
                _export_epub(source_path, temporary, parsed, translations)
            else:
                _export_epub_plain(temporary, parsed, translations)
        elif suffix == ".pdf":
            layout = parsed.metadata.get("pdfLayout")
            if parsed.format == "pdf" and layout == "book":
                _export_pdf_layout(source_path, temporary, parsed, translations)
            elif parsed.format == "pdf" and layout == "scan":
                _export_pdf_scan(source_path, temporary, parsed, translations)
            else:
                _export_pdf_text(temporary, parsed, translations)
        os.replace(temporary, destination)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def save_parsed(path: Path, parsed: ParsedDocument) -> None:
    path.write_text(
        json.dumps(
            {
                "title": parsed.title,
                "format": parsed.format,
                "warnings": parsed.warnings,
                "metadata": parsed.metadata,
                "blocks": [block_to_dict(block) for block in parsed.blocks],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def load_parsed(path: Path) -> ParsedDocument:
    data = json.loads(path.read_text(encoding="utf-8"))
    return ParsedDocument(
        data["title"],
        data["format"],
        tuple(block_from_dict(item) for item in data["blocks"]),
        tuple(data.get("warnings", [])),
        data.get("metadata", {}),
    )
